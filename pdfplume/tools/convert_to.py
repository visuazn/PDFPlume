"""Convert to PDF: images, scans, Word, PowerPoint, Excel, HTML."""

import csv
import html
import io
import os
import re
import tempfile
import urllib.parse
import urllib.request
from pathlib import Path

import pymupdf as fitz
from PIL import Image, ImageFilter, ImageOps

from .. import engines
from ..utils import parse_color, to_bool
from . import ToolError, tool

PAGE_SIZES = {"a4": fitz.paper_rect("a4"), "letter": fitz.paper_rect("letter"),
              "legal": fitz.paper_rect("legal"), "a3": fitz.paper_rect("a3"),
              "a5": fitz.paper_rect("a5")}
MARGINS = {"none": 0, "small": 20, "big": 50}


# ----------------------------------------------------------------------- images
def _image_frames(path):
    """Yield PIL images for every frame of ``path`` (multi-page TIFF support)."""
    im = Image.open(path)
    frames = getattr(im, "n_frames", 1) if (im.format or "").upper() == "TIFF" else 1
    for i in range(frames):
        im.seek(i)
        yield ImageOps.exif_transpose(im.copy())


def _image_bytes(im, filt="color"):
    """Encode a PIL image for embedding; JPEG for photos, PNG for the rest."""
    if filt == "grayscale":
        im = ImageOps.grayscale(im.convert("RGB"))
    elif filt == "bw":
        g = ImageOps.autocontrast(ImageOps.grayscale(im.convert("RGB")), cutoff=2)
        g = g.filter(ImageFilter.MedianFilter(3))
        im = g.point(lambda x: 255 if x > 150 else 0, mode="1")
    elif filt == "enhance":
        rgb = im.convert("RGB")
        rgb = ImageOps.autocontrast(rgb, cutoff=1)
        im = rgb.filter(ImageFilter.UnsharpMask(radius=2, percent=80, threshold=3))
    buf = io.BytesIO()
    if im.mode in ("RGBA", "LA", "P", "1"):
        if im.mode == "P":
            im = im.convert("RGBA")
        im.save(buf, "PNG", optimize=True)
    else:
        if im.mode not in ("RGB", "L", "CMYK"):
            im = im.convert("RGB")
        im.save(buf, "JPEG", quality=92, optimize=True)
    return buf.getvalue()


def images_to_pdf(doc, paths, page_size="a4", orientation="portrait", margin="small",
                  filt="color", progress=None, crops=None):
    """``crops[n]`` per image: None (keep), "auto" (find the page) or four corners."""
    m = MARGINS.get(margin, 20)
    for n, path in enumerate(paths):
        if progress:
            progress(n / max(len(paths), 1), f"Adding image {n + 1}")
        try:
            frames = list(_image_frames(path))
        except Exception:
            raise ToolError(f"'{Path(path).name}' is not a supported image.")
        crop = crops[n] if crops else None
        for im in frames:
            if crop is not None:
                from .. import scan
                corners = scan.detect_document(im)["corners"] if crop == "auto" else crop
                im = scan.warp(im, corners)
            data = _image_bytes(im, filt)
            iw, ih = im.size
            if page_size == "fit":
                # 96 dpi is how images are usually meant to be shown on screen.
                w, h = iw * 72 / 96, ih * 72 / 96
                scale = min(1.0, 14000 / max(w + 2 * m, h + 2 * m))
                page = doc.new_page(width=w * scale + 2 * m, height=h * scale + 2 * m)
            else:
                base = PAGE_SIZES.get(page_size, PAGE_SIZES["a4"])
                landscape = orientation == "landscape" or (orientation == "auto" and iw > ih)
                pw, ph = (base.height, base.width) if landscape else (base.width, base.height)
                page = doc.new_page(width=pw, height=ph)
            area = page.rect + (m, m, -m, -m)
            page.insert_image(area, stream=data, keep_proportion=True)
    return doc


@tool("jpg_to_pdf", accept=("image",), zip_name="pdfplume_images.zip")
def jpg_to_pdf(ctx):
    size = ctx.opt("page_size", "a4")
    orient = ctx.opt("orientation", "portrait")
    margin = ctx.opt("margin", "none")
    if to_bool(ctx.opt("merge"), True):
        doc = fitz.open()
        images_to_pdf(doc, [i.path for i in ctx.inputs], size, orient, margin,
                      progress=ctx.progress)
        name = ctx.inputs[0].stem if len(ctx.inputs) == 1 else "images"
        ctx.save_pdf(doc, f"{name}.pdf")
    else:
        for inp in ctx.inputs:
            doc = fitz.open()
            images_to_pdf(doc, [inp.path], size, orient, margin)
            ctx.save_pdf(doc, f"{inp.stem}.pdf")


@tool("scan_to_pdf", accept=("image",))
def scan_to_pdf(ctx):
    auto = to_bool(ctx.opt("auto_crop"), True)
    crops = []
    for inp in ctx.inputs:
        corners = inp.options.get("corners")
        if inp.options.get("crop") is False:
            crops.append(None)
        elif isinstance(corners, list) and len(corners) == 4:
            crops.append(corners)
        else:
            crops.append("auto" if auto else None)
    doc = fitz.open()
    images_to_pdf(doc, [i.path for i in ctx.inputs], ctx.opt("page_size", "a4"),
                  "auto", ctx.opt("margin", "none"), filt=ctx.opt("filter", "enhance"),
                  progress=ctx.progress, crops=crops)
    ctx.save_pdf(doc, "scan.pdf")


# ------------------------------------------------------------------- html engine
BASE_CSS = """
body { font-family: sans-serif; font-size: 11pt; line-height: 1.35; color: #111; }
h1 { font-size: 20pt; } h2 { font-size: 16pt; } h3 { font-size: 13pt; }
p { margin: 0 0 6pt 0; }
table { border-collapse: collapse; margin: 4pt 0; }
td, th { border: 0.5pt solid #999; padding: 2pt 4pt; vertical-align: top; }
th { background: #eee; font-weight: bold; }
pre, code { font-family: monospace; font-size: 9pt; }
"""


def render_html(sections, out_path, page_size="a4", landscape=False, margin=40, css=BASE_CSS,
                archive=None):
    """Lay out one or more HTML documents into a PDF using MuPDF's Story engine.

    Every section starts on a new page.
    """
    if isinstance(sections, str):
        sections = [sections]
    mediabox = PAGE_SIZES.get(page_size, PAGE_SIZES["a4"])
    if landscape:
        mediabox = fitz.Rect(0, 0, mediabox.height, mediabox.width)
    where = mediabox + (margin, margin, -margin, -margin)
    writer = fitz.DocumentWriter(str(out_path))
    pages = 0
    for section in sections:
        story = fitz.Story(html=section, user_css=css, archive=archive)
        more = True
        guard = 0
        while more and guard < 5000:
            dev = writer.begin_page(mediabox)
            more, _ = story.place(where)
            story.draw(dev)
            writer.end_page()
            pages += 1
            guard += 1
    if pages == 0:
        dev = writer.begin_page(mediabox)
        writer.end_page()
    writer.close()
    return out_path


# ------------------------------------------------------------- office fallbacks
def _docx_to_html(path):
    import mammoth

    def convert_image(image):
        with image.open() as f:
            data = f.read()
        attrs = {"src": f"data:{image.content_type};base64,"
                        + __import__("base64").b64encode(data).decode()}
        try:
            with Image.open(io.BytesIO(data)) as im:
                w, h = im.size
            max_w = 500
            if w > max_w:
                h = h * max_w / w
                w = max_w
            attrs["width"] = str(int(w))
            attrs["height"] = str(int(h))
        except Exception:
            pass
        return attrs

    with open(path, "rb") as f:
        result = mammoth.convert_to_html(f, convert_image=mammoth.images.img_element(convert_image))
    return f"<html><body>{result.value}</body></html>"


def _text_to_html(path):
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    paras = "".join(f"<p>{html.escape(line) or '&#160;'}</p>" for line in text.splitlines())
    return f"<html><body>{paras}</body></html>"


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:,.2f}".rstrip("0").rstrip(".") if value != int(value) else f"{int(value):,}"
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d") if not getattr(value, "hour", 0) else str(value)
    return str(value)


def _rows_to_table(rows):
    out = ["<table>"]
    for r, row in enumerate(rows):
        tag = "th" if r == 0 else "td"
        out.append("<tr>" + "".join(f"<{tag}>{html.escape(_cell_text(v))}</{tag}>" for v in row)
                   + "</tr>")
    out.append("</table>")
    return "".join(out)


def _spreadsheet_sections(path):
    ext = Path(path).suffix.lower()
    sections, widest = [], 0
    if ext == ".csv":
        raw = Path(path).read_text(encoding="utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(raw[:4096])
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(raw), dialect))
        widest = max((len(r) for r in rows), default=0)
        sections.append(f"<html><body>{_rows_to_table(rows)}</body></html>")
        return sections, widest
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        # Trim empty trailing rows / columns.
        while rows and all(v is None for v in rows[-1]):
            rows.pop()
        if not rows:
            continue
        width = max((max((i + 1 for i, v in enumerate(r) if v is not None), default=0)
                     for r in rows), default=0)
        rows = [r[:width] for r in rows]
        widest = max(widest, width)
        sections.append(f"<html><body><h3>{html.escape(ws.title)}</h3>"
                        f"{_rows_to_table(rows)}</body></html>")
    return sections or ["<html><body><p>(empty workbook)</p></body></html>"], widest


EMU_PER_PT = 12700


def _pptx_to_pdf(path, out_path):
    """Basic PowerPoint renderer used when LibreOffice is not available."""
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation(path)
    sw, sh = prs.slide_width / EMU_PER_PT, prs.slide_height / EMU_PER_PT
    doc = fitz.open()

    def rect_of(shape):
        return fitz.Rect(Emu(shape.left or 0).pt, Emu(shape.top or 0).pt,
                         Emu((shape.left or 0) + (shape.width or 0)).pt,
                         Emu((shape.top or 0) + (shape.height or 0)).pt)

    def color_of(font_or_fill):
        try:
            rgb = font_or_fill.color.rgb if hasattr(font_or_fill, "color") else font_or_fill.fore_color.rgb
            return f"#{rgb}" if rgb else None
        except Exception:
            return None

    def text_html(shape, default_size):
        parts = []
        for para in shape.text_frame.paragraphs:
            align = {1: "left", 2: "center", 3: "right", 4: "justify"}.get(
                int(para.alignment) if para.alignment is not None else 1, "left")
            runs = []
            for run in para.runs:
                f = run.font
                style = [f"font-size:{(f.size.pt if f.size else default_size):.1f}px"]
                if f.bold:
                    style.append("font-weight:bold")
                if f.italic:
                    style.append("font-style:italic")
                c = color_of(f)
                if c:
                    style.append(f"color:{c}")
                runs.append(f"<span style=\"{';'.join(style)}\">{html.escape(run.text)}</span>")
            bullet = "• " if para.level and runs else ""
            parts.append(f"<p style=\"text-align:{align};margin:0 0 4px {para.level * 18}px\">"
                         f"{bullet}{''.join(runs) or '&#160;'}</p>")
        return "".join(parts)

    def draw_shapes(page, shapes):
        for shape in shapes:
            try:
                if shape.shape_type == 6:  # group
                    draw_shapes(page, shape.shapes)
                    continue
                r = rect_of(shape)
                if r.is_empty:
                    continue
                if hasattr(shape, "fill"):
                    try:
                        if shape.fill.type == 1:
                            c = color_of(shape.fill)
                            if c:
                                page.draw_rect(r, color=None, fill=parse_color(c))
                    except Exception:
                        pass
                image = None
                try:
                    image = shape.image  # pictures and picture placeholders
                except (AttributeError, ValueError):
                    pass
                if image is not None:
                    page.insert_image(r, stream=image.blob, keep_proportion=False)
                if getattr(shape, "has_table", False) and shape.has_table:
                    rows = [[c.text for c in row.cells] for row in shape.table.rows]
                    page.insert_htmlbox(r, _rows_to_table(rows), css=BASE_CSS + "*{font-size:10px}")
                elif getattr(shape, "has_text_frame", False) and shape.has_text_frame \
                        and shape.text_frame.text.strip():
                    is_title = getattr(shape, "is_placeholder", False) and \
                        shape.placeholder_format.type in (1, 3)
                    page.insert_htmlbox(r + (0, 0, 0, r.height), text_html(shape, 36 if is_title else 18),
                                        css="*{font-family:sans-serif}")
            except Exception:
                continue

    for slide in prs.slides:
        page = doc.new_page(width=sw, height=sh)
        try:
            bg = slide.background.fill
            if bg.type == 1:
                page.draw_rect(page.rect, color=None, fill=parse_color(f"#{bg.fore_color.rgb}"))
        except Exception:
            pass
        draw_shapes(page, slide.shapes)
    if doc.page_count == 0:
        doc.new_page(width=sw, height=sh)
    doc.save(out_path, garbage=3, deflate=True)


def office_to_pdf(ctx, inp, kind):
    """Convert one office document; LibreOffice first, pure Python fallback second."""
    if engines.soffice_path():
        out_dir = ctx.work_dir / inp.id if inp.id else ctx.work_dir
        src = out_dir / f"{inp.stem}{inp.ext}"
        out_dir.mkdir(parents=True, exist_ok=True)
        src.write_bytes(inp.path.read_bytes())
        try:
            return engines.soffice_convert(src, out_dir / "pdf")
        except engines.EngineError:
            if inp.ext not in (".docx", ".xlsx", ".csv", ".pptx", ".txt"):
                raise
    out = ctx.work_path(f"{inp.stem}.pdf")
    ext = inp.ext
    if kind == "word" and ext == ".docx":
        render_html(_docx_to_html(inp.path), out)
    elif kind == "word" and ext == ".txt":
        render_html(_text_to_html(inp.path), out)
    elif kind == "excel" and ext in (".xlsx", ".csv", ".xlsm"):
        sections, widest = _spreadsheet_sections(inp.path)
        render_html(sections, out, landscape=widest > 6, margin=28,
                    css=BASE_CSS + "td,th{font-size:8pt}")
    elif kind == "powerpoint" and ext == ".pptx":
        _pptx_to_pdf(inp.path, out)
    else:
        raise ToolError(f"Converting {ext} files needs LibreOffice, which is not installed. "
                        "Install LibreOffice (free) and restart PDFPlume.")
    return out


def _office_tool(ctx, kind):
    import shutil
    n = len(ctx.inputs)
    for i, inp in enumerate(ctx.inputs):
        ctx.progress(i / n, f"Converting {inp.name}", force=True)
        pdf = office_to_pdf(ctx, inp, kind)
        out = ctx.output_path(f"{inp.stem}.pdf")
        shutil.copyfile(pdf, out)
        ctx.add_output(out)
    ctx.extra["engine"] = "libreoffice" if engines.soffice_path() else "basic"


@tool("word_to_pdf", accept=("word",), zip_name="pdfplume_word.zip")
def word_to_pdf(ctx):
    _office_tool(ctx, "word")


@tool("powerpoint_to_pdf", accept=("powerpoint",), zip_name="pdfplume_powerpoint.zip")
def powerpoint_to_pdf(ctx):
    _office_tool(ctx, "powerpoint")


@tool("excel_to_pdf", accept=("excel",), zip_name="pdfplume_excel.zip")
def excel_to_pdf(ctx):
    _office_tool(ctx, "excel")


# --------------------------------------------------------------------- html
VIEWPORTS = {"desktop": 1440, "laptop": 1280, "tablet": 768, "mobile": 390}
PAPER_INCHES = {"a4": (8.27, 11.69), "letter": (8.5, 11), "legal": (8.5, 14), "a3": (11.69, 16.54)}
MARGIN_INCHES = {"none": 0, "small": 0.4, "big": 0.8}


def _html_playwright(url, out, opts):
    from playwright.sync_api import sync_playwright
    vw = VIEWPORTS.get(opts["viewport"], 1440)
    with sync_playwright() as p:
        try:
            browser = p.chromium.launch()
        except Exception:
            exe = engines.chrome_path()
            if not exe:
                raise
            browser = p.chromium.launch(executable_path=exe)
        try:
            page = browser.new_page(viewport={"width": vw, "height": 900})
            try:
                page.goto(url, wait_until="networkidle", timeout=45000)
            except Exception:
                page.goto(url, wait_until="load", timeout=60000)
            page.emulate_media(media="screen")
            # Trigger lazy loaded images.
            page.evaluate("""async () => { for (let y = 0; y < document.body.scrollHeight;
                y += 700) { window.scrollTo(0, y); await new Promise(r => setTimeout(r, 60)); }
                window.scrollTo(0, 0); }""")
            if opts["page_size"] == "one_page":
                height = page.evaluate("Math.max(document.body.scrollHeight,"
                                       "document.documentElement.scrollHeight)")
                page.pdf(path=str(out), width=f"{vw}px", height=f"{min(height + 20, 19000)}px",
                         print_background=True, page_ranges="1")
            else:
                pw, ph = PAPER_INCHES.get(opts["page_size"], PAPER_INCHES["a4"])
                if opts["landscape"]:
                    pw, ph = ph, pw
                m = MARGIN_INCHES.get(opts["margin"], 0.4)
                scale = max(0.1, min(2.0, (pw - 2 * m) * 96 / vw))
                page.pdf(path=str(out), width=f"{pw}in", height=f"{ph}in", scale=scale,
                         margin={k: f"{m}in" for k in ("top", "bottom", "left", "right")},
                         print_background=True)
        finally:
            browser.close()


def _html_chrome_cli(url, out, opts):
    exe = engines.chrome_path()
    vw = VIEWPORTS.get(opts["viewport"], 1440)
    with tempfile.TemporaryDirectory(prefix="pdfplume_chrome_") as profile:
        cmd = [exe, "--disable-gpu", "--no-first-run", "--hide-scrollbars",
               "--disable-background-networking", "--disable-component-update",
               "--disable-extensions", "--disable-sync", "--no-default-browser-check",
               "--disable-dev-shm-usage",
               f"--user-data-dir={profile}", "--no-pdf-header-footer",
               "--run-all-compositor-stages-before-draw", "--virtual-time-budget=15000",
               f"--window-size={vw},1000", f"--print-to-pdf={out}", url]
        if "headless-shell" not in Path(exe).name:
            cmd.insert(1, "--headless=new")
        if hasattr(os, "geteuid") and os.geteuid() == 0:
            cmd.insert(1, "--no-sandbox")
        engines.run_program(cmd, 90, "Chrome", check=False)
    if not Path(out).exists() or Path(out).stat().st_size == 0:
        raise ToolError("Chrome could not print this page.")


def _html_basic(url, out, opts, work_dir):
    """Last resort: fetch the HTML and lay it out with MuPDF (no JavaScript)."""
    tmp = Path(tempfile.mkdtemp(dir=work_dir))
    if url.startswith("file://"):
        raw = Path(urllib.request.url2pathname(urllib.parse.urlparse(url).path)).read_bytes()
    else:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 PDFPlume"})
        try:
            with urllib.request.urlopen(req, timeout=30) as resp:
                raw = resp.read()
                url = resp.url
        except Exception as exc:
            raise ToolError(f"Could not load the page: {exc}")
    text = raw.decode("utf-8", "replace")
    text = re.sub(r"(?is)<(script|style|noscript|iframe|svg)[^>]*>.*?</\1>", "", text)
    text = re.sub(r"(?is)<link[^>]*>", "", text)
    count = 0

    def fetch_img(m):
        nonlocal count
        src = html.unescape(m.group(2))
        if count >= 40 or src.startswith("data:") or url.startswith("file://"):
            return m.group(0)
        count += 1
        try:
            full = urllib.parse.urljoin(url, src)
            req = urllib.request.Request(full, headers={"User-Agent": "Mozilla/5.0 PDFPlume"})
            with urllib.request.urlopen(req, timeout=10) as resp:
                data = resp.read(8 * 1024 * 1024)
            with Image.open(io.BytesIO(data)) as im:
                im = im.convert("RGB")
                if im.width > 500:
                    im = im.resize((500, int(im.height * 500 / im.width)))
                name = f"img{count}.jpg"
                im.save(tmp / name, "JPEG", quality=85)
            return f'{m.group(1)}"{name}"'
        except Exception:
            return f'{m.group(1)}""'

    text = re.sub(r'(?i)(<img[^>]*?\ssrc=)["\']([^"\']+)["\']', fetch_img, text)
    m = MARGIN_INCHES.get(opts["margin"], 0.4) * 72
    size = opts["page_size"] if opts["page_size"] in PAGE_SIZES else "a4"
    render_html(text, out, size, opts["landscape"], margin=m or 10,
                css=BASE_CSS + "img{max-width:100%}", archive=fitz.Archive(str(tmp)))


@tool("html_to_pdf", accept=("html",), min_files=0)
def html_to_pdf(ctx):
    opts = {
        "viewport": ctx.opt("viewport", "desktop"),
        "page_size": ctx.opt("page_size", "a4"),
        "landscape": ctx.opt("orientation", "portrait") == "landscape",
        "margin": ctx.opt("margin", "small"),
    }
    targets = []
    urls = [u for u in (ctx.opt("urls") or []) if str(u).strip()] or [ctx.opt("url") or ""]
    for url in urls:
        url = str(url).strip()
        if not url:
            continue
        if not re.match(r"^[a-z]+://", url, re.I):
            url = "https://" + url
        name = re.sub(r"[^\w.-]+", "_", urllib.parse.urlparse(url).netloc or "page")
        targets.append((url, name))
    for inp in ctx.inputs:
        targets.append((inp.path.resolve().as_uri(), inp.stem))
    if not targets:
        raise ToolError("Enter a website URL or add an HTML file.")
    engine_used = None
    for i, (target, name) in enumerate(targets):
        ctx.progress(i / len(targets), f"Rendering {target[:60]}", force=True)
        out = ctx.output_path(f"{name}.pdf")
        errors = []
        for engine, fn in (("playwright", _html_playwright), ("chrome", _html_chrome_cli)):
            if engine == "playwright" and not engines.playwright_available():
                continue
            if engine == "chrome" and not engines.chrome_path():
                continue
            try:
                fn(target, out, opts)
                engine_used = engine
                break
            except Exception as exc:
                errors.append(str(exc))
        else:
            _html_basic(target, out, opts, ctx.work_dir)
            engine_used = "basic"
        ctx.add_output(out)
    ctx.extra["engine"] = engine_used

