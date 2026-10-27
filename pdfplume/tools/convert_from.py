"""Convert from PDF: JPG, Word, PowerPoint, Excel, PDF/A."""

import io
import logging
import re

import pymupdf as fitz

from .. import engines
from ..utils import to_bool
from . import ToolError, tool
from .optimize import ocr_document, page_needs_ocr


# ------------------------------------------------------------------------- JPG
def _save_pixmap(pix, path, fmt):
    if fmt == "jpg":
        pix.save(str(path), jpg_quality=90)
    else:
        pix.save(str(path))


@tool("pdf_to_jpg", zip_name="pdfplume_images.zip")
def pdf_to_jpg(ctx):
    mode = ctx.opt("mode", "pages")
    dpi = {"normal": 150, "high": 300, "low": 96}.get(ctx.opt("quality", "normal"), 150)
    fmt = "png" if ctx.opt("format") == "png" else "jpg"
    for inp in ctx.inputs:
        doc = ctx.open_pdf(inp)
        if mode == "extract":
            seen = set()
            count = 0
            for page in doc:
                ctx.progress(page.number / doc.page_count, "Extracting images")
                for img in page.get_images(full=True):
                    xref, smask = img[0], img[1]
                    if xref in seen:
                        continue
                    seen.add(xref)
                    try:
                        pix = fitz.Pixmap(doc, xref)
                        if pix.width < 16 or pix.height < 16:
                            continue
                        if smask:
                            try:
                                pix = fitz.Pixmap(pix, fitz.Pixmap(doc, smask))
                            except Exception:
                                pass
                        if pix.colorspace and pix.colorspace.n > 3:
                            pix = fitz.Pixmap(fitz.csRGB, pix)
                        if fmt == "jpg" and pix.alpha:
                            pix = fitz.Pixmap(pix, 0)
                        count += 1
                        out = ctx.output_path(f"{inp.stem}_image_{count}.{fmt}")
                        _save_pixmap(pix, out, fmt)
                        ctx.add_output(out)
                    except Exception:
                        continue
            if count == 0 and len(ctx.inputs) == 1:
                raise ToolError("No images were found in this PDF.")
        else:
            for page in doc:
                ctx.progress(page.number / doc.page_count, f"Converting page {page.number + 1}")
                pix = page.get_pixmap(dpi=dpi, alpha=False)
                out = ctx.output_path(f"{inp.stem}_page_{page.number + 1}.{fmt}")
                _save_pixmap(pix, out, fmt)
                ctx.add_output(out)


def _prepare_source(ctx, inp, ocr=False):
    """Return a path to an unencrypted (and optionally OCR'd) copy of the input."""
    src = ctx.decrypted_path(inp)
    if ocr and engines.tesseract_path():
        doc = fitz.open(src)
        if any(page_needs_ocr(p) for p in doc):
            ocr_document(doc, ctx.opt("language", "eng"), work_dir=ctx.work_dir,
                         progress=lambda f, m: ctx.progress(f * 0.4, m))
            src = ctx.work_path(f"{inp.stem}_ocr.pdf")
            doc.save(src, garbage=3, deflate=True)
        doc.close()
    return src


# ------------------------------------------------------------------------ Word
def _simple_docx(pdf_path, out):
    import docx
    from docx.shared import Pt
    document = docx.Document()
    doc = fitz.open(pdf_path)
    for page in doc:
        for block in page.get_text("blocks", sort=True):
            if block[6] != 0:
                continue
            text = re.sub(r"\s*\n\s*", " ", block[4]).strip()
            if text:
                document.add_paragraph(text).style.font.size = Pt(11)
        if page.number < doc.page_count - 1:
            document.add_page_break()
    document.save(out)


@tool("pdf_to_word", zip_name="pdfplume_word.zip")
def pdf_to_word(ctx):
    ocr = to_bool(ctx.opt("ocr"))
    n = len(ctx.inputs)
    for i, inp in enumerate(ctx.inputs):
        ctx.progress(i / n, f"Converting {inp.name}", force=True)
        src = _prepare_source(ctx, inp, ocr)
        out = ctx.output_path(f"{inp.stem}.docx")
        try:
            from pdf2docx import Converter
            logging.getLogger().setLevel(logging.WARNING)
            cv = Converter(str(src))
            try:
                cv.convert(str(out), multi_processing=False)
            finally:
                cv.close()
        except Exception:
            _simple_docx(src, out)
        ctx.add_output(out)


# ------------------------------------------------------------------ PowerPoint
def _font_name(name):
    base = name.split("+")[-1]
    low = base.lower()
    if "times" in low or "serif" in low and "sans" not in low:
        return "Times New Roman"
    if "courier" in low or "mono" in low:
        return "Courier New"
    if "helvetica" in low or "arial" in low:
        return "Arial"
    return re.split(r"[-,]", base)[0] or "Arial"


@tool("pdf_to_powerpoint", zip_name="pdfplume_powerpoint.zip")
def pdf_to_powerpoint(ctx):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Emu, Pt

    editable = ctx.opt("mode", "editable") == "editable"
    n = len(ctx.inputs)
    for fi, inp in enumerate(ctx.inputs):
        src_path = _prepare_source(ctx, inp, to_bool(ctx.opt("ocr")))
        doc = fitz.open(src_path)
        first = doc[0].rect
        prs = Presentation()
        prs.slide_width = Emu(int(first.width * 12700))
        prs.slide_height = Emu(int(first.height * 12700))
        blank = prs.slide_layouts[6]
        for page in doc:
            ctx.progress((fi + page.number / doc.page_count) / n,
                         f"Converting page {page.number + 1}")
            tmp = fitz.open()
            tmp.insert_pdf(doc, from_page=page.number, to_page=page.number)
            tp = tmp[0]
            if tp.rotation:
                tp.remove_rotation()
            scale = min(first.width / tp.rect.width, first.height / tp.rect.height)
            ox = (first.width - tp.rect.width * scale) / 2
            oy = (first.height - tp.rect.height * scale) / 2
            lines = []
            if editable:
                data = tp.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE)
                for block in data["blocks"]:
                    if block.get("type") != 0:
                        continue
                    for line in block["lines"]:
                        if abs(line["dir"][1]) > 0.01:
                            continue  # rotated text stays in the background image
                        if any(s["text"].strip() for s in line["spans"]):
                            lines.append(line)
                if lines:
                    for line in lines:
                        tp.add_redact_annot(fitz.Rect(line["bbox"]))
                    tp.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE,
                                        graphics=fitz.PDF_REDACT_LINE_ART_NONE,
                                        text=fitz.PDF_REDACT_TEXT_REMOVE)
            pix = tp.get_pixmap(dpi=150, alpha=False)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")),
                                     Emu(int(ox * 12700)), Emu(int(oy * 12700)),
                                     width=Emu(int(tp.rect.width * scale * 12700)),
                                     height=Emu(int(tp.rect.height * scale * 12700)))
            for line in lines:
                x0, y0, x1, y1 = line["bbox"]
                box = slide.shapes.add_textbox(
                    Emu(int((ox + x0 * scale) * 12700)), Emu(int((oy + y0 * scale) * 12700)),
                    Emu(int(max(x1 - x0, 4) * scale * 12700 * 1.06)),
                    Emu(int(max(y1 - y0, 4) * scale * 12700)))
                tf = box.text_frame
                tf.word_wrap = False
                tf.auto_size = MSO_AUTO_SIZE.NONE
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                para = tf.paragraphs[0]
                for span in line["spans"]:
                    if not span["text"]:
                        continue
                    run = para.add_run()
                    run.text = span["text"]
                    font = run.font
                    font.size = Pt(max(1, span["size"] * scale))
                    font.name = _font_name(span["font"])
                    font.bold = bool(span["flags"] & 16) or "bold" in span["font"].lower()
                    font.italic = bool(span["flags"] & 2) or "italic" in span["font"].lower()
                    c = span.get("color", 0)
                    font.color.rgb = RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255)
            tmp.close()
        out = ctx.output_path(f"{inp.stem}.pptx")
        prs.save(out)
        ctx.add_output(out)


# ----------------------------------------------------------------------- Excel
NUM_RE = re.compile(r"^[-+]?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?$")


def _cell_value(text):
    if text is None:
        return None
    t = str(text).strip()
    if NUM_RE.match(t):
        v = float(t.replace(",", ""))
        return int(v) if v.is_integer() and "." not in t else v
    return t


def _text_rows(page):
    """Rebuild a grid from plain text lines by splitting on wide gaps."""
    words = page.get_text("words", sort=True)
    lines = {}
    for w in words:
        key = (w[5], w[6])
        lines.setdefault(key, []).append(w)
    rows = []
    for key in sorted(lines, key=lambda k: (lines[k][0][1], lines[k][0][0])):
        ws = sorted(lines[key], key=lambda w: w[0])
        cells, current, last_x1 = [], [], None
        for w in ws:
            height = w[3] - w[1]
            if last_x1 is not None and w[0] - last_x1 > max(height * 1.2, 8):
                cells.append(" ".join(current))
                current = []
            current.append(w[4])
            last_x1 = w[2]
        if current:
            cells.append(" ".join(current))
        rows.append(cells)
    return rows


@tool("pdf_to_excel", zip_name="pdfplume_excel.zip")
def pdf_to_excel(ctx):
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    single = ctx.opt("layout", "per_page") == "single"
    n = len(ctx.inputs)
    for fi, inp in enumerate(ctx.inputs):
        src = _prepare_source(ctx, inp, to_bool(ctx.opt("ocr")))
        doc = fitz.open(src)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        sheet = wb.create_sheet("Data") if single else None
        tables_found = 0
        for page in doc:
            ctx.progress((fi + page.number / doc.page_count) / n,
                         f"Reading tables on page {page.number + 1}")
            blocks = []
            try:
                found = page.find_tables()
                for t in found.tables:
                    # Read each cell ourselves: Table.extract() can scramble punctuation.
                    rows = [[page.get_textbox(fitz.Rect(c)).strip() if c else ""
                             for c in row.cells] for row in t.rows]
                    if rows and any(any(c for c in r) for r in rows):
                        blocks.append(rows)
            except Exception:
                pass
            tables_found += len(blocks)
            if not blocks:
                rows = _text_rows(page)
                if rows:
                    blocks.append(rows)
            if not single:
                sheet = wb.create_sheet(f"Page {page.number + 1}")
            for rows in blocks:
                start = sheet.max_row + (2 if sheet.max_row > 1 else 0)
                if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
                    start = 1
                for r, row in enumerate(rows):
                    for c, value in enumerate(row):
                        cell = sheet.cell(start + r, c + 1, _cell_value(
                            value.replace("\n", " ") if isinstance(value, str) else value))
                        if r == 0 and tables_found:
                            cell.font = Font(bold=True)
        for ws in wb.worksheets:
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 8), 60)
        if not wb.worksheets:
            wb.create_sheet("Data")
        out = ctx.output_path(f"{inp.stem}.xlsx")
        wb.save(out)
        ctx.add_output(out)
        ctx.extra["tables_found"] = ctx.extra.get("tables_found", 0) + tables_found


# ----------------------------------------------------------------- text / md
def _plain_text(doc, page_markers=True):
    parts = []
    for page in doc:
        text = page.get_text("text", sort=True).strip()
        parts.append(f"--- Page {page.number + 1} ---\n{text}" if page_markers else text)
    return "\n\n".join(parts).strip() + "\n"


def load_layout():
    """Turn on PyMuPDF's layout model (better reading order and tables) if it is installed."""
    import sys
    if "pymupdf.layout" in sys.modules and sys.modules["pymupdf.layout"] is None:
        del sys.modules["pymupdf.layout"]  # blocked at start-up, see pdfplume/__init__.py
    import pymupdf4llm
    try:
        pymupdf4llm.use_layout(True)
    except Exception:
        pymupdf4llm.use_layout(False)
    return pymupdf4llm


def _markdown(doc):
    import contextlib
    try:
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            return load_layout().to_markdown(doc, show_progress=False, use_ocr=False)
    except Exception:
        return _plain_text(doc, page_markers=False)


@tool("pdf_to_text", zip_name="pdfplume_text.zip")
def pdf_to_text(ctx):
    fmt = "txt" if ctx.opt("format") == "txt" else "md"
    n = len(ctx.inputs)
    for i, inp in enumerate(ctx.inputs):
        ctx.progress(i / n, f"Reading {inp.name}", force=True)
        src = _prepare_source(ctx, inp, to_bool(ctx.opt("ocr")))
        with fitz.open(src) as doc:
            text = _markdown(doc) if fmt == "md" else _plain_text(doc, to_bool(ctx.opt("page_markers"), True))
        out = ctx.output_path(f"{inp.stem}.{fmt}")
        out.write_text(text, encoding="utf-8")
        ctx.add_output(out)
        if i == 0:
            ctx.extra["preview"] = text[:8000]
            ctx.extra["words"] = len(text.split())
    ctx.extra["format"] = fmt


# ----------------------------------------------------------------------- PDF/A
def _srgb_icc():
    try:
        from PIL import ImageCms
        return ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()
    except Exception:
        import glob
        for path in glob.glob("/usr/share/color/icc/ghostscript/srgb.icc"):
            return open(path, "rb").read()
    return None


def _finish_pdfa(path, out, part, conformance):
    """Add the XMP identification and sRGB output intent PDF/A requires."""
    import pikepdf
    with pikepdf.open(path) as pdf:
        icc = _srgb_icc()
        if icc and "/OutputIntents" not in pdf.Root:
            stream = pikepdf.Stream(pdf, icc)
            stream["/N"] = 3
            intent = pikepdf.Dictionary(
                Type=pikepdf.Name.OutputIntent, S=pikepdf.Name.GTS_PDFA1,
                OutputConditionIdentifier=pikepdf.String("sRGB IEC61966-2.1"),
                Info=pikepdf.String("sRGB IEC61966-2.1"), DestOutputProfile=stream)
            pdf.Root.OutputIntents = pikepdf.Array([pdf.make_indirect(intent)])
        for key in ("/OpenAction", "/AA"):
            if key in pdf.Root:
                del pdf.Root[key]
        names = pdf.Root.get("/Names")
        if names is not None and "/JavaScript" in names:
            del names["/JavaScript"]
        with pdf.open_metadata(set_pikepdf_as_editor=False) as meta:
            meta["pdfaid:part"] = str(part)
            meta["pdfaid:conformance"] = conformance
            if not meta.get("dc:title"):
                meta["dc:title"] = str(pdf.docinfo.get("/Title", "") or "Untitled")
        pdf.save(out, fix_metadata_version=True, static_id=False)


@tool("pdf_to_pdfa", zip_name="pdfplume_pdfa.zip")
def pdf_to_pdfa(ctx):
    level = str(ctx.opt("conformance", "2b")).lower()
    part = level[0] if level[:1] in "123" else "2"
    conformance = (level[1:] or "b").upper()
    if conformance not in ("A", "B", "U"):
        conformance = "B"
    for inp in ctx.inputs:
        src = ctx.decrypted_path(inp)
        tmp = ctx.work_path(f"{inp.stem}_tmp.pdf")
        engine = "basic"
        if engines.ghostscript_path():
            try:
                engines.ghostscript([
                    f"-dPDFA={part}", "-sDEVICE=pdfwrite", "-sColorConversionStrategy=RGB",
                    "-dPDFACompatibilityPolicy=1", "-dAutoRotatePages=/None",
                    "-dEmbedAllFonts=true", f"-sOutputFile={tmp}", str(src)])
                engine = "ghostscript"
            except engines.EngineError:
                pass
        if engine == "basic":
            doc = fitz.open(src)
            doc.save(tmp, garbage=3, deflate=True)
            doc.close()
        out = ctx.output_path(f"{inp.stem}_pdfa.pdf")
        _finish_pdfa(tmp, out, part, conformance)
        ctx.add_output(out)
        ctx.extra["engine"] = engine
    ctx.extra["conformance"] = f"PDF/A-{part}{conformance.lower()}"
    if not engines.ghostscript_path():
        ctx.extra["note"] = ("Install Ghostscript for full PDF/A conversion (font embedding "
                             "and colour conversion).")

