import io
import os
import shutil
import sys
from pathlib import Path

import pymupdf as fitz
import pytest
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ["PDFPLUME_OLLAMA_AUTODETECT"] = "0"  # tests must not depend on a local Ollama

from pdfplume.storage import Storage  # noqa: E402
from pdfplume.tools import registry, run_job  # noqa: E402
from pdfplume.utils import file_kind  # noqa: E402

LOREM = ("PDF tools make documents easier to manage. Merging joins files together. "
         "Splitting separates pages into new documents. Compression reduces the file size "
         "while keeping the quality acceptable. Security tools protect sensitive content. "
         "Contact us at hello@example.com or call +1 555 123 4567 for details. ")


def make_sample_pdf(path, pages=5, text_prefix="Page"):
    doc = fitz.open()
    img = Image.new("RGB", (400, 300), (30, 120, 220))
    ImageDraw.Draw(img).ellipse((50, 50, 350, 250), fill=(250, 200, 0))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=95)
    for i in range(pages):
        page = doc.new_page()
        page.insert_text((72, 72), f"{text_prefix} {i + 1} heading", fontsize=22)
        page.insert_textbox(fitz.Rect(72, 100, 520, 300), LOREM * 2, fontsize=11)
        if i == 0:
            page.insert_image(fitz.Rect(72, 320, 472, 620), stream=buf.getvalue())
        if i == 1:
            # A simple ruled table.
            x0, y0, cw, rh = 72, 340, 120, 24
            rows = [["Name", "Qty", "Price"], ["Apple", "3", "1.50"], ["Pear", "10", "2,450.00"]]
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    rect = fitz.Rect(x0 + c * cw, y0 + r * rh, x0 + (c + 1) * cw, y0 + (r + 1) * rh)
                    page.draw_rect(rect, color=(0, 0, 0), width=0.8)
                    page.insert_text((rect.x0 + 4, rect.y1 - 7), val, fontsize=11)
    doc.set_metadata({"title": "Sample", "author": "Tester"})
    doc.save(path)
    return path


DESK_PAGE = [[400, 100], [1150, 160], [1210, 1120], [340, 1080]]  # paper corners in the photo


def make_desk_photo(path):
    """A photo of a sheet of paper lying askew on a dark, noisy desk."""
    import cv2
    import numpy as np
    rng = np.random.default_rng(3)
    w, h = 1600, 1200
    desk = (rng.normal(0, 18, (h, w, 3)) + np.array([70, 50, 35])).clip(0, 255).astype(np.uint8)
    desk = cv2.GaussianBlur(desk, (7, 7), 0)
    sheet = Image.new("RGB", (850, 1100), (245, 243, 238))
    draw = ImageDraw.Draw(sheet)
    for i in range(30):
        draw.rectangle((80, 90 + i * 32, 80 + int(rng.integers(300, 690)), 104 + i * 32), fill=(30, 30, 30))
    m = cv2.getPerspectiveTransform(np.float32([[0, 0], [849, 0], [849, 1099], [0, 1099]]), np.float32(DESK_PAGE))
    warped = cv2.warpPerspective(np.asarray(sheet), m, (w, h))
    mask = cv2.warpPerspective(np.full((1100, 850), 255, np.uint8), m, (w, h))
    photo = np.where(mask[..., None] > 0, warped, desk)
    photo = (photo * np.linspace(0.75, 1.05, w)[None, :, None]).clip(0, 255).astype(np.uint8)
    Image.fromarray(photo).save(path, quality=88)
    return path


def make_form_pdf(path):
    """A two-page form: text, multi-line text, check box, combo, list, radio group, read-only."""
    doc = fitz.open()
    page = doc.new_page()

    def add(pg, kind, name, rect, **kw):
        w = fitz.Widget()
        w.field_type, w.field_name, w.rect = kind, name, fitz.Rect(rect)
        for k, v in kw.items():
            setattr(w, k, v)
        pg.add_widget(w)

    page.insert_text((72, 60), "Application form", fontsize=18)
    add(page, fitz.PDF_WIDGET_TYPE_TEXT, "full_name", (72, 80, 300, 100), field_label="Full name")
    add(page, fitz.PDF_WIDGET_TYPE_TEXT, "notes", (72, 110, 400, 170), field_flags=1 << 12)
    add(page, fitz.PDF_WIDGET_TYPE_CHECKBOX, "agree", (72, 180, 86, 194))
    add(page, fitz.PDF_WIDGET_TYPE_COMBOBOX, "country", (72, 204, 220, 222), choice_values=["India", "Germany", "USA"])
    add(page, fitz.PDF_WIDGET_TYPE_LISTBOX, "colour", (72, 232, 220, 280), choice_values=["Red", "Green", "Blue"])
    add(page, fitz.PDF_WIDGET_TYPE_RADIOBUTTON, "size", (72, 290, 86, 304), field_value=True)
    add(page, fitz.PDF_WIDGET_TYPE_RADIOBUTTON, "size", (120, 290, 134, 304), field_value=True)
    add(page, fitz.PDF_WIDGET_TYPE_TEXT, "ref", (72, 320, 300, 340), field_value="FIXED-1", field_flags=1)
    page2 = doc.new_page()
    page2.set_rotation(90)
    add(page2, fitz.PDF_WIDGET_TYPE_TEXT, "city", (72, 80, 300, 100))
    # Give the radio buttons realistic 'on' state names, both off.
    page = doc[0]  # adding a page invalidates earlier page objects
    radios = [w for w in page.widgets() if w.field_type == fitz.PDF_WIDGET_TYPE_RADIOBUTTON]
    for w, name in zip(radios, ("Small", "Medium")):
        ap = doc.xref_get_key(w.xref, "AP/N")[1].replace("/True", f"/{name}")
        doc.xref_set_key(w.xref, "AP/N", ap)
        doc.xref_set_key(w.xref, "AS", "/Off")
    doc.save(path)
    return path


@pytest.fixture(scope="session")
def samples(tmp_path_factory):
    d = tmp_path_factory.mktemp("samples")
    files = {}
    files["pdf"] = make_sample_pdf(d / "sample.pdf")
    files["pdf2"] = make_sample_pdf(d / "other.pdf", pages=3, text_prefix="Section")

    # A modified copy for compare.
    doc = fitz.open(files["pdf"])
    page = doc[0]
    page.add_redact_annot(page.search_for("Merging joins")[0])
    page.apply_redactions()
    page.insert_text((72, 700), "A brand new sentence was added here.", fontsize=11)
    doc.save(d / "modified.pdf")
    files["modified"] = d / "modified.pdf"

    rot = fitz.open(files["pdf"])
    rot[0].set_rotation(90)
    rot.save(d / "rotated.pdf")
    files["rotated"] = d / "rotated.pdf"

    enc = fitz.open(files["pdf"])
    enc.save(d / "locked.pdf", encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="secret",
             owner_pw="owner")
    files["locked"] = d / "locked.pdf"

    # "Scanned" document: text only as pixels.
    src = fitz.open(files["pdf"])
    scan = fitz.open()
    for p in src.pages(0, 2):
        pix = p.get_pixmap(dpi=200)
        np_ = scan.new_page(width=p.rect.width, height=p.rect.height)
        np_.insert_image(np_.rect, pixmap=pix)
    scan.save(d / "scanned.pdf")
    files["scanned"] = d / "scanned.pdf"

    im = Image.new("RGB", (800, 600), (200, 30, 30))
    ImageDraw.Draw(im).rectangle((100, 100, 700, 500), fill=(255, 255, 255))
    im.save(d / "photo.jpg", "JPEG")
    files["jpg"] = d / "photo.jpg"
    from pillow_heif import register_heif_opener
    register_heif_opener()
    im.save(d / "iphone.heic", "HEIF")
    files["heic"] = d / "iphone.heic"

    # A photo-heavy PDF (noise does not compress) for size-targeted compression.
    import random
    rnd = random.Random(7)
    noisy = Image.frombytes("RGB", (1600, 1200), bytes(rnd.getrandbits(8) for _ in range(1600 * 1200 * 3)))
    buf = io.BytesIO()
    noisy.save(buf, "JPEG", quality=95)
    photo = fitz.open()
    for _ in range(2):
        photo.new_page().insert_image(fitz.Rect(36, 36, 559, 428), stream=buf.getvalue())
    photo.save(d / "photos.pdf")
    files["photos"] = d / "photos.pdf"

    files["form"] = make_form_pdf(d / "form.pdf")
    files["desk"] = make_desk_photo(d / "desk.jpg")

    rgba = Image.new("RGBA", (300, 500), (0, 128, 0, 128))
    rgba.save(d / "logo.png")
    files["png"] = d / "logo.png"

    import docx
    document = docx.Document()
    document.add_heading("Quarterly report", 1)
    document.add_paragraph("Revenue grew strongly this quarter. " * 10)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Region", "Sales"
    table.cell(1, 0).text, table.cell(1, 1).text = "North", "1200"
    document.save(d / "report.docx")
    files["docx"] = d / "report.docx"

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Product", "Units", "Price"])
    for i in range(30):
        ws.append([f"Item {i}", i * 3, i * 1.25])
    wb.save(d / "sheet.xlsx")
    files["xlsx"] = d / "sheet.xlsx"
    (d / "data.csv").write_text("a,b,c\n1,2,3\n4,5,6\n")
    files["csv"] = d / "data.csv"

    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Hello slides"
    slide.placeholders[1].text = "First bullet\nSecond bullet"
    slide.shapes.add_picture(str(files["jpg"]), Inches(5), Inches(4), width=Inches(3))
    prs.save(d / "deck.pptx")
    files["pptx"] = d / "deck.pptx"

    (d / "page.html").write_text("<html><body><h1>Hello HTML</h1><p>Some <b>bold</b> text."
                                 "</p></body></html>")
    files["html"] = d / "page.html"
    (d / "notes.txt").write_text("Plain text line one\nLine two")
    files["txt"] = d / "notes.txt"
    return files


@pytest.fixture()
def storage(tmp_path):
    return Storage(tmp_path / "data")


@pytest.fixture()
def run_tool(storage):
    """Upload the given files and run a tool in-process. Returns (result, out_dir)."""
    registry()

    def run(slug, paths, options=None, passwords=None, file_options=None):
        inputs = []
        for i, p in enumerate(paths):
            p = Path(p)
            meta = storage.add_file(p.name, lambda dst, src=p: shutil.copyfile(src, dst))
            meta = storage.get_meta(meta["id"])
            assert meta["kind"] == file_kind(p.suffix)
            if passwords and passwords[i]:
                meta["password"] = passwords[i]
            if file_options:
                meta["options"] = file_options[i]
            inputs.append(meta)
        job_id, job_dir = storage.new_job()
        result = run_job(slug, inputs, options or {}, str(job_dir))
        return result, job_dir / "out"

    return run
