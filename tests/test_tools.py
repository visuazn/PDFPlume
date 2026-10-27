"""End-to-end tests for every tool handler."""

import base64
import io
import zipfile

import pymupdf as fitz
import pytest
from PIL import Image

from pdfplume import engines
from pdfplume.utils import PageRangeError, parse_page_ranges


def ok(result):
    assert result["status"] == "done", result.get("message")
    return result


def pdf(path):
    return fitz.open(path)


def zip_names(path):
    with zipfile.ZipFile(path) as zf:
        return zf.namelist()


needs_tesseract = pytest.mark.skipif(not engines.tesseract_path(), reason="Tesseract missing")
needs_gs = pytest.mark.skipif(not engines.ghostscript_path(), reason="Ghostscript missing")
needs_soffice = pytest.mark.skipif(not engines.soffice_path(), reason="LibreOffice missing")


# ------------------------------------------------------------------ helpers
def test_page_ranges():
    assert parse_page_ranges("1-3,5", 6) == [0, 1, 2, 4]
    assert parse_page_ranges("4-", 6) == [3, 4, 5]
    assert parse_page_ranges("-2", 6) == [0, 1]
    assert parse_page_ranges("all", 3) == [0, 1, 2]
    assert parse_page_ranges("3-1", 3) == [2, 1, 0]
    assert parse_page_ranges("last", 4) == [3]
    with pytest.raises(PageRangeError):
        parse_page_ranges("9", 3)
    with pytest.raises(PageRangeError):
        parse_page_ranges("a-b", 3)


# ----------------------------------------------------------------- organize
def test_merge(run_tool, samples):
    res, out = run_tool("merge_pdf", [samples["pdf"], samples["pdf2"]],
                        file_options=[{}, {"rotation": 90}])
    ok(res)
    doc = pdf(out / res["file"])
    assert doc.page_count == 8
    assert doc[5].rotation == 90 and doc[0].rotation == 0


def test_merge_needs_two_files(run_tool, samples):
    from pdfplume.tools import TOOLS
    assert TOOLS["merge_pdf"]["min_files"] == 2


def test_merge_locked_requires_password(run_tool, samples):
    res, _ = run_tool("merge_pdf", [samples["pdf"], samples["locked"]])
    assert res["status"] == "error" and "password" in res["message"].lower()
    res, out = run_tool("merge_pdf", [samples["pdf"], samples["locked"]], passwords=[None, "secret"])
    assert pdf(out / ok(res)["file"]).page_count == 10


@pytest.mark.parametrize("options,count", [
    ({"mode": "ranges", "ranges": [{"from": 1, "to": 2}, {"from": 4, "to": 5}]}, 2),
    ({"mode": "ranges", "ranges": [{"from": 1, "to": 2}, {"from": 4, "to": 5}], "merge": True}, 1),
    ({"mode": "fixed", "every": 2}, 3),
    ({"mode": "all"}, 5),
    ({"mode": "select", "pages": "1,3", "merge": False}, 2),
    ({"mode": "select", "pages": "1,3"}, 1),
    ({"mode": "size", "max_mb": 0.05}, None),
])
def test_split(run_tool, samples, options, count):
    res, out = run_tool("split_pdf", [samples["pdf"]], options)
    ok(res)
    if count is not None:
        assert res["count"] == count
    if res["count"] > 1:
        assert res["file"].endswith(".zip")
        assert len(zip_names(out / res["file"])) == res["count"]


def test_split_bad_range(run_tool, samples):
    res, _ = run_tool("split_pdf", [samples["pdf"]], {"mode": "ranges", "ranges": [{"from": 4, "to": 9}]})
    assert res["status"] == "error"


def test_remove_and_extract(run_tool, samples):
    res, out = run_tool("remove_pages", [samples["pdf"]], {"pages": "2,4-5"})
    doc = pdf(out / ok(res)["file"])
    assert doc.page_count == 2 and "Page 3" in doc[1].get_text()
    res, _ = run_tool("remove_pages", [samples["pdf"]], {"pages": "1-5"})
    assert res["status"] == "error"
    res, out = run_tool("extract_pages", [samples["pdf"]], {"pages": "5,1"})
    doc = pdf(out / ok(res)["file"])
    assert doc.page_count == 2 and "Page 5" in doc[0].get_text()
    res, out = run_tool("extract_pages", [samples["pdf"]], {"pages": "1-3", "separate": True})
    assert ok(res)["count"] == 3


def test_organize(run_tool, samples):
    layout = [{"file": 1, "page": 0}, {"file": 0, "page": 2, "rotation": 180},
              {"page": None, "blank": True}, {"file": 0, "page": 0}]
    res, out = run_tool("organize_pdf", [samples["pdf"], samples["pdf2"]], {"pages": layout})
    doc = pdf(out / ok(res)["file"])
    assert doc.page_count == 4
    assert "Section 1" in doc[0].get_text()
    assert doc[1].rotation == 180
    assert doc[2].get_text().strip() == ""


def test_rotate(run_tool, samples):
    res, out = run_tool("rotate_pdf", [samples["pdf"]], {},
                        file_options=[{"rotation": 270, "pages": {"1": 90}}])
    doc = pdf(out / ok(res)["file"])
    assert doc[0].rotation == 270 and doc[1].rotation == 0


# ----------------------------------------------------------------- optimize
@pytest.mark.parametrize("level", ["extreme", "recommended", "low"])
def test_compress(run_tool, samples, level):
    res, out = run_tool("compress_pdf", [samples["pdf"]], {"level": level})
    ok(res)
    assert res["extra"]["after"] <= res["extra"]["before"]
    assert pdf(out / res["file"]).page_count == 5


def test_compress_fallback_without_ghostscript(run_tool, samples, monkeypatch):
    monkeypatch.setattr(engines, "ghostscript_path", lambda: None)
    res, _ = run_tool("compress_pdf", [samples["pdf"]], {"level": "extreme"})
    assert ok(res)["extra"]["engine"] == "mupdf"


def test_repair(run_tool, samples, tmp_path):
    broken = tmp_path / "broken.pdf"
    data = samples["pdf"].read_bytes()
    idx = data.rfind(b"xref")
    broken.write_bytes(data[:idx] + b"garbage" + data[idx + 40:])
    res, out = run_tool("repair_pdf", [broken])
    assert pdf(out / ok(res)["file"]).page_count == 5


@needs_tesseract
def test_ocr(run_tool, samples):
    res, out = run_tool("ocr_pdf", [samples["scanned"]], {"language": "eng"})
    ok(res)
    assert res["extra"]["pages_recognized"] == 2
    text = pdf(out / res["file"])[0].get_text()
    assert "heading" in text.lower()


# ------------------------------------------------------------- convert to
def test_jpg_to_pdf(run_tool, samples):
    res, out = run_tool("jpg_to_pdf", [samples["jpg"], samples["png"]],
                        {"page_size": "a4", "orientation": "auto", "margin": "small"})
    doc = pdf(out / ok(res)["file"])
    assert doc.page_count == 2
    assert doc[0].rect.width > doc[0].rect.height  # landscape photo
    res, out = run_tool("jpg_to_pdf", [samples["jpg"], samples["png"]],
                        {"page_size": "fit", "merge": False})
    assert ok(res)["count"] == 2


def test_scan_to_pdf(run_tool, samples):
    for f in ("bw", "grayscale", "enhance", "color"):
        res, out = run_tool("scan_to_pdf", [samples["jpg"]], {"filter": f})
        assert pdf(out / ok(res)["file"]).page_count == 1


@pytest.mark.parametrize("slug,key", [("word_to_pdf", "docx"), ("word_to_pdf", "txt"),
                                      ("excel_to_pdf", "xlsx"), ("excel_to_pdf", "csv"),
                                      ("powerpoint_to_pdf", "pptx")])
def test_office_fallback(run_tool, samples, monkeypatch, slug, key):
    monkeypatch.setattr(engines, "soffice_path", lambda: None)
    res, out = run_tool(slug, [samples[key]])
    doc = pdf(out / ok(res)["file"])
    assert doc.page_count >= 1
    assert doc[0].get_text().strip()


@needs_soffice
@pytest.mark.parametrize("slug,key", [("word_to_pdf", "docx"), ("excel_to_pdf", "xlsx"),
                                      ("powerpoint_to_pdf", "pptx")])
def test_office_libreoffice(run_tool, samples, slug, key):
    res, out = run_tool(slug, [samples[key]])
    assert ok(res)["extra"]["engine"] == "libreoffice"
    assert pdf(out / res["file"])[0].get_text().strip()


def test_html_file_basic(run_tool, samples, monkeypatch):
    monkeypatch.setattr(engines, "playwright_available", lambda: False)
    monkeypatch.setattr(engines, "chrome_path", lambda: None)
    res, out = run_tool("html_to_pdf", [samples["html"]])
    assert ok(res)["extra"]["engine"] == "basic"
    assert "Hello HTML" in pdf(out / res["file"])[0].get_text()


@pytest.mark.skipif(not engines.chrome_path(), reason="Chrome missing")
def test_html_file_chrome(run_tool, samples, monkeypatch):
    monkeypatch.setattr(engines, "playwright_available", lambda: False)
    res, out = run_tool("html_to_pdf", [samples["html"]])
    assert ok(res)["extra"]["engine"] == "chrome"
    assert "Hello HTML" in pdf(out / res["file"])[0].get_text()


def test_html_requires_input(run_tool):
    res, _ = run_tool("html_to_pdf", [], {})
    assert res["status"] == "error"


# ----------------------------------------------------------- convert from
def test_pdf_to_jpg(run_tool, samples):
    res, out = run_tool("pdf_to_jpg", [samples["pdf"]], {"mode": "pages"})
    assert ok(res)["count"] == 5 and res["file"].endswith(".zip")
    res, out = run_tool("pdf_to_jpg", [samples["pdf"]], {"mode": "extract"})
    ok(res)
    assert res["count"] == 1 and res["file"].endswith(".jpg")
    with Image.open(out / res["file"]) as im:
        assert im.size == (400, 300)


def test_pdf_to_word(run_tool, samples):
    import docx
    res, out = run_tool("pdf_to_word", [samples["pdf"]])
    text = "\n".join(p.text for p in docx.Document(out / ok(res)["file"]).paragraphs)
    assert "Page 1 heading" in text


def test_pdf_to_powerpoint(run_tool, samples):
    from pptx import Presentation
    res, out = run_tool("pdf_to_powerpoint", [samples["rotated"]])
    prs = Presentation(out / ok(res)["file"])
    assert len(prs.slides) == 5
    texts = [sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame]
    assert any("Page 2 heading" in t for t in texts)


def test_pdf_to_excel(run_tool, samples):
    import openpyxl
    res, out = run_tool("pdf_to_excel", [samples["pdf"]])
    wb = openpyxl.load_workbook(out / ok(res)["file"])
    assert len(wb.worksheets) == 5
    values = [c.value for row in wb.worksheets[1].iter_rows() for c in row]
    assert "Apple" in values and 2450.0 in values
    assert res["extra"]["tables_found"] >= 1


def test_pdf_to_pdfa(run_tool, samples, monkeypatch):
    import pikepdf
    res, out = run_tool("pdf_to_pdfa", [samples["pdf"]], {"conformance": "2b"})
    with pikepdf.open(out / ok(res)["file"]) as p:
        assert "/OutputIntents" in p.Root
        with p.open_metadata() as m:
            assert m["pdfaid:part"] == "2"
    monkeypatch.setattr(engines, "ghostscript_path", lambda: None)
    res, out = run_tool("pdf_to_pdfa", [samples["pdf"]], {"conformance": "1b"})
    assert ok(res)["extra"]["engine"] == "basic"


# --------------------------------------------------------------------- edit
def test_page_numbers(run_tool, samples):
    res, out = run_tool("add_page_numbers", [samples["rotated"]],
                        {"format": "Page {n} of {p}", "first_number": 3, "pages": "2-5",
                         "position": "bottom-right"})
    doc = pdf(out / ok(res)["file"])
    assert "Page 3 of 6" in doc[1].get_text()
    assert "Page 6 of 6" in doc[4].get_text()
    assert "of" not in doc[0].get_text().split("heading")[0]


@pytest.mark.parametrize("opts", [
    {"type": "text", "text": "CONFIDENTIAL", "rotation": 45, "opacity": 40},
    {"type": "text", "text": "DRAFT", "mosaic": True, "layer": "below", "bold": True},
    {"type": "text", "text": "Top", "position": "top-left", "underline": True, "pages": "1-2"},
])
def test_watermark_text(run_tool, samples, opts):
    res, out = run_tool("add_watermark", [samples["rotated"]], opts)
    doc = pdf(out / ok(res)["file"])
    assert opts["text"] in doc[0].get_text()


def test_watermark_image(run_tool, samples, storage):
    res, out = run_tool("add_watermark", [samples["pdf"]],
                        {"type": "image", "image_id_path": str(samples["png"]), "rotation": 30})
    doc = pdf(out / ok(res)["file"])
    assert len(doc[2].get_images()) >= 1


def test_crop(run_tool, samples):
    box = {"x0": 0.1, "y0": 0.1, "x1": 0.6, "y1": 0.5}
    res, out = run_tool("crop_pdf", [samples["pdf"]], {"box": box})
    doc = pdf(out / ok(res)["file"])
    r = doc[0].rect
    assert abs(r.width - 595 * 0.5) < 1 and abs(r.height - 842 * 0.4) < 1
    res, out = run_tool("crop_pdf", [samples["rotated"]], {"box": box, "apply": "current", "page": 0})
    doc = pdf(out / ok(res)["file"])
    assert abs(doc[0].rect.width - 842 * 0.5) < 1  # rotated page: visible width is 842
    assert doc[1].rect.width == 595


def _png_data_url():
    im = Image.new("RGBA", (120, 40), (0, 0, 200, 255))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


EDIT_ELEMENTS = [
    {"type": "text", "page": 0, "x": 0.1, "y": 0.8, "w": 0.4, "h": 0.03, "text": "Added text Ünïcödé",
     "size": 14, "color": "#ff0000", "bold": True},
    {"type": "image", "page": 0, "x": 0.5, "y": 0.8, "w": 0.2, "h": 0.05, "src": _png_data_url(),
     "opacity": 0.5},
    {"type": "rect", "page": 1, "x": 0.1, "y": 0.1, "w": 0.2, "h": 0.1, "stroke": "#000", "fill": "#00ff00"},
    {"type": "ellipse", "page": 1, "x": 0.4, "y": 0.1, "w": 0.2, "h": 0.1, "stroke": "#00f"},
    {"type": "line", "page": 1, "x1": 0.1, "y1": 0.5, "x2": 0.9, "y2": 0.5, "stroke": "#f00", "width": 3},
    {"type": "arrow", "page": 1, "x1": 0.1, "y1": 0.6, "x2": 0.5, "y2": 0.7, "stroke": "#f00", "width": 2},
    {"type": "path", "page": 1, "points": [[0.1, 0.8], [0.2, 0.85], [0.3, 0.8]], "stroke": "#000", "width": 2},
    {"type": "highlight", "page": 2, "x": 0.1, "y": 0.1, "w": 0.5, "h": 0.03},
    {"type": "whiteout", "page": 2, "x": 0.1, "y": 0.2, "w": 0.5, "h": 0.05},
]


def test_edit(run_tool, samples):
    res, out = run_tool("edit_pdf", [samples["rotated"]], {"elements": EDIT_ELEMENTS})
    doc = pdf(out / ok(res)["file"])
    assert "Added text Ünïcödé" in doc[0].get_text()
    assert doc[0].rotation == 0
    assert len(doc[1].get_drawings()) > 5
    res, _ = run_tool("edit_pdf", [samples["pdf"]], {"elements": []})
    assert res["status"] == "error"


# ----------------------------------------------------------------- security
def test_protect_and_unlock(run_tool, samples):
    res, out = run_tool("protect_pdf", [samples["pdf"]], {"password": "pw123", "allow_copy": False})
    locked = out / ok(res)["file"]
    doc = pdf(locked)
    assert doc.needs_pass and doc.authenticate("pw123")
    res, out = run_tool("unlock_pdf", [locked])
    assert res["status"] == "error"
    res, out = run_tool("unlock_pdf", [locked], passwords=["pw123"])
    doc = pdf(out / ok(res)["file"])
    assert not doc.needs_pass and not doc.is_encrypted and doc.page_count == 5


def test_protect_mismatch(run_tool, samples):
    res, _ = run_tool("protect_pdf", [samples["pdf"]], {"password": "a", "confirm": "b"})
    assert res["status"] == "error"


def test_sign(run_tool, samples):
    elements = [{"type": "image", "page": 4, "x": 0.6, "y": 0.85, "w": 0.25, "h": 0.06,
                 "src": _png_data_url()},
                {"type": "text", "page": 4, "x": 0.6, "y": 0.92, "w": 0.3, "h": 0.02,
                 "text": "John Doe 2026-10-27", "size": 10}]
    res, out = run_tool("sign_pdf", [samples["pdf"]], {"elements": elements})
    doc = pdf(out / ok(res)["file"])
    assert "John Doe" in doc[4].get_text() and doc[4].get_images()


def test_redact(run_tool, samples):
    res, out = run_tool("redact_pdf", [samples["rotated"]],
                        {"terms": ["Merging"], "patterns": ["email", "phone"],
                         "areas": [{"page": 1, "x0": 0, "y0": 0, "x1": 1, "y1": 0.12}]})
    doc = pdf(out / ok(res)["file"])
    text = "".join(p.get_text() for p in doc)
    assert "Merging" not in text
    assert "hello@example.com" not in text
    assert "555" not in text
    assert "Page 2 heading" not in doc[1].get_text()
    assert "Page 3 heading" in doc[2].get_text()
    assert doc.metadata.get("author") in ("", None)


def test_redact_search_preview(samples):
    from pdfplume.tools.security import find_matches
    doc = fitz.open(samples["rotated"])
    matches = find_matches(doc, ["PDF tools"], [])
    assert len(matches) == 10
    first = [m for m in matches if m["page"] == 0][0]["frac"]
    # page 0 is rotated 90°, so the hit appears near the right edge.
    assert first[0] > 0.5
    second = [m for m in matches if m["page"] == 1][0]["frac"]
    assert second[0] < 0.3


def test_compare_text(run_tool, samples):
    res, out = run_tool("compare_pdf", [samples["pdf"], samples["modified"]])
    ok(res)
    extra = res["extra"]
    assert extra["mode"] == "text" and not extra["identical"]
    kinds = {c["kind"] for c in extra["changes"]}
    assert "ins" in kinds or "chg" in kinds
    assert any("brand new sentence" in c["b"] for c in extra["changes"])
    assert pdf(out / res["file"]).page_count == 6


def test_compare_overlay(run_tool, samples):
    res, out = run_tool("compare_pdf", [samples["pdf"], samples["modified"]], {"mode": "overlay"})
    ok(res)
    pages = res["extra"]["pages"]
    assert pages[0]["changed"] > 0 and pages[1]["changed"] == 0
    assert (out / "assets" / pages[0]["asset"]).exists()


# ------------------------------------------------------------- intelligence
def test_summarize(run_tool, samples):
    res, out = run_tool("summarize_pdf", [samples["pdf"]], {"length": "short"})
    ok(res)
    x = res["extra"]
    assert x["summary"] and len(x["points"]) == len(x["summary"])
    assert all(1 <= pt["page"] <= 5 and pt["text"] for pt in x["points"])
    assert x["stats"]["pages"] == 5
    assert [s["title"] for s in x["sections"]] == [f"Page {i} heading" for i in range(1, 6)]
    assert [s["page"] for s in x["sections"]] == [1, 2, 3, 4, 5]
    assert all(pt["page"] == s["page"] for s in x["sections"] for pt in s["points"])
    assert all(s["points"] for s in x["sections"])
    text = pdf(out / res["file"])[0].get_text()
    assert text.startswith("Summary") and "(p. 1)" in text and "Section by section" in pdf(out / res["file"]).get_page_text(-1) + text


def test_summarize_without_sections(run_tool, samples):
    res, _ = run_tool("summarize_pdf", [samples["pdf"]], {"sections": False})
    assert ok(res)["extra"]["sections"] == []


def test_summarize_with_ollama(run_tool, samples, monkeypatch):
    from pdfplume.tools import intelligence
    seen = {}

    def fake(prompt, timeout=600):
        seen["prompt"] = prompt
        return "A short summary.\n- Compression reduces size (p. 1)"
    monkeypatch.setattr(intelligence, "_ollama", fake)
    res, _ = run_tool("summarize_pdf", [samples["pdf"]], {"engine": "ollama"})
    assert ok(res)["extra"]["engine"] == "ollama" and "(p. 1)" in res["extra"]["ai_summary"]
    assert "[Page 1]" in seen["prompt"] and "[Page 5]" in seen["prompt"]


def test_translate_without_engine(run_tool, samples):
    res, _ = run_tool("translate_pdf", [samples["pdf"]], {"target": "es"})
    assert res["status"] == "error" and "translation engine" in res["message"].lower()


def test_translate_with_fake_engine(run_tool, samples, monkeypatch):
    from pdfplume.tools import intelligence
    monkeypatch.setattr(intelligence, "translation_status",
                        lambda: {"available": True, "engines": ["fake"], "languages": []})
    monkeypatch.setattr(intelligence.Translator, "translate_many",
                        lambda self, texts: [t.upper() for t in texts])
    res, out = run_tool("translate_pdf", [samples["pdf"]], {"target": "es"})
    text = pdf(out / ok(res)["file"])[0].get_text()
    assert "PAGE 1 HEADING" in text and "Page 1 heading" not in text


# ---------------------------------------------------------------- quick wins
def test_heic_to_pdf(run_tool, samples):
    res, out = run_tool("jpg_to_pdf", [samples["heic"]], {"page_size": "fit"})
    doc = pdf(out / ok(res)["file"])
    assert doc.page_count == 1 and doc[0].get_images()


def test_compress_to_target(run_tool, samples):
    before = samples["photos"].stat().st_size
    assert before > 1_000_000
    res, out = run_tool("compress_pdf", [samples["photos"]], {"mode": "target", "target_mb": 0.3})
    ok(res)
    assert (out / res["file"]).stat().st_size <= 0.3 * 1024 * 1024
    assert "under" in res["extra"]["note"]
    # A target larger than the file keeps it as it is.
    res, out = run_tool("compress_pdf", [samples["pdf"]], {"mode": "target", "target_mb": 50})
    assert ok(res)["extra"]["saved_percent"] == 0


def test_compress_target_unreachable(run_tool, samples, tmp_path):
    import os
    big = fitz.open()
    for _ in range(20):  # distinct images: identical ones would be stored once
        im = Image.frombytes("RGB", (800, 600), os.urandom(800 * 600 * 3))
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=90)
        big.new_page().insert_image(fitz.Rect(36, 36, 559, 428), stream=buf.getvalue())
    big.save(tmp_path / "album.pdf")
    res, out = run_tool("compress_pdf", [tmp_path / "album.pdf"], {"mode": "target", "target_mb": 0.05})
    assert "smallest version" in ok(res)["extra"]["note"]
    assert (out / res["file"]).stat().st_size < (tmp_path / "album.pdf").stat().st_size


@pytest.mark.parametrize("fmt,ext", [("md", ".md"), ("txt", ".txt")])
def test_pdf_to_text(run_tool, samples, fmt, ext):
    res, out = run_tool("pdf_to_text", [samples["pdf"]], {"format": fmt})
    ok(res)
    assert res["file"].endswith(ext)
    text = (out / res["file"]).read_text(encoding="utf-8")
    assert "Page 1 heading" in text and "Page 5 heading" in text
    assert res["extra"]["preview"].startswith(text[:50])
    if fmt == "md":
        assert "# Page 1 heading" in text


def test_cancel_flag_stops_job(storage, samples):
    import shutil
    from pdfplume.tools import run_job
    meta = storage.add_file("s.pdf", lambda p: shutil.copyfile(samples["pdf"], p))
    job_id, job_dir = storage.new_job()
    (job_dir / "cancel").write_text("1")
    result = run_job("compress_pdf", [storage.get_meta(meta["id"])], {}, str(job_dir))
    assert result["status"] == "cancelled"


def test_cancel_kills_external_program():
    import sys
    import time
    from pdfplume.engines import Cancelled, run_program, set_cancel_check
    set_cancel_check(lambda: True)
    started = time.time()
    try:
        with pytest.raises(Cancelled):
            run_program([sys.executable, "-c", "import time; time.sleep(30)"], 60, "sleeper")
    finally:
        set_cancel_check(None)
    assert time.time() - started < 5


# ------------------------------------------------------------------ layout
@pytest.mark.parametrize("per,sheets", [(2, 3), (4, 2), (9, 1), (16, 1)])
def test_nup(run_tool, samples, per, sheets):
    res, out = run_tool("nup_pdf", [samples["pdf"]], {"per_sheet": per, "border": True})
    doc = pdf(out / ok(res)["file"])
    assert doc.page_count == sheets
    assert "Page 1 heading" in doc[0].get_text()
    if per == 2:
        assert doc[0].rect.width > doc[0].rect.height  # two portrait pages fit best side by side


def test_booklet_order():
    from pdfplume.tools.layout import booklet_order
    assert booklet_order(8) == [(7, 0), (1, 6), (5, 2), (3, 4)]
    assert booklet_order(5)[0] == (None, 0)  # padded with blanks


def test_booklet(run_tool, samples):
    res, out = run_tool("booklet_pdf", [samples["pdf"]], {})
    doc = pdf(out / ok(res)["file"])
    assert doc.page_count == 4 and res["extra"]["sheets"] == 2  # 5 pages -> 8 -> 2 sheets
    words = doc[0].get_text("words")
    right = " ".join(w[4] for w in words if w[0] > doc[0].rect.width / 2)
    assert "Page 1" in right  # first sheet: blank | page 1


def test_resize(run_tool, samples):
    letter = fitz.open()
    letter.new_page(width=612, height=792).insert_text((72, 72), "Letter page", fontsize=20)
    letter.save(samples["pdf"].parent / "letter.pdf")
    res, out = run_tool("resize_pdf", [samples["pdf"].parent / "letter.pdf"], {"size": "a4", "margin": 10})
    doc = pdf(out / ok(res)["file"])
    assert abs(doc[0].rect.width - 595.3) < 1 and abs(doc[0].rect.height - 841.9) < 1
    assert "Letter page" in doc[0].get_text()
    res, out = run_tool("resize_pdf", [samples["pdf"]], {"size": "custom", "width_mm": 100, "height_mm": 100})
    assert abs(pdf(out / ok(res)["file"])[0].rect.width - 283.5) < 1


def test_grayscale(run_tool, samples):
    import numpy as np
    res, out = run_tool("grayscale_pdf", [samples["pdf"]])
    doc = pdf(out / ok(res)["file"])
    pix = doc[0].get_pixmap(dpi=30)
    a = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.h, pix.w, pix.n).astype(int)
    assert (a.max(axis=2) - a.min(axis=2)).max() <= 2
    assert "Page 1 heading" in doc[0].get_text()


# ------------------------------------------------------------------ forms
def _fields(path):
    """Fields by name; radio buttons as size#0, size#1 (xrefs change when a PDF is saved)."""
    from pdfplume.tools.forms import list_fields
    out, seen = {}, {}
    for f in list_fields(pdf(path)):
        key = f["name"]
        if f["type"] == "radio":
            seen[key] = seen.get(key, -1) + 1
            key = f"{key}#{seen[key]}"
        out[key] = f
    return out


def test_list_fields(samples):
    from pdfplume.tools.forms import list_fields
    fields = list_fields(pdf(samples["form"]))
    by = {f["name"]: f for f in fields}
    assert by["full_name"]["label"] == "Full name" and by["full_name"]["type"] == "text"
    assert by["notes"]["multiline"] and by["ref"]["readonly"] and by["ref"]["value"] == "FIXED-1"
    assert [o["value"] for o in by["country"]["options"]] == ["India", "Germany", "USA"]
    assert by["colour"]["type"] == "listbox"
    radios = [f for f in fields if f["type"] == "radio"]
    assert len(radios) == 2 and not any(r["value"] for r in radios)
    x0, y0, x1, y1 = by["full_name"]["rect"]
    assert 0 < x0 < x1 < 1 and 0 < y0 < y1 < 1
    city = by["city"]  # on a page rotated by 90 degrees: the field runs top to bottom
    assert city["page"] == 1 and (city["rect"][3] - city["rect"][1]) > (city["rect"][2] - city["rect"][0])


def test_fill_form(run_tool, samples):
    fields = _fields(samples["form"])
    radio = ["size#0", "size#1"]
    values = {
        fields["full_name"]["xref"]: "Ada Lovelace",
        fields["notes"]["xref"]: "Line one\nLine two",
        fields["agree"]["xref"]: True,
        fields["country"]["xref"]: "Germany",
        fields["colour"]["xref"]: "Blue",
        fields[radio[0]]["xref"]: False,
        fields[radio[1]]["xref"]: True,
        fields["ref"]["xref"]: "CHANGED",
        fields["city"]["xref"]: "Kochi",
    }
    res, out = run_tool("fill_form", [samples["form"]], {"values": {str(k): v for k, v in values.items()}})
    path = out / ok(res)["file"]
    after = _fields(path)
    assert after["full_name"]["value"] == "Ada Lovelace"
    assert after["notes"]["value"] == "Line one\nLine two"
    assert after["agree"]["value"] is True
    assert after["country"]["value"] == "Germany" and after["colour"]["value"] == "Blue"
    assert after["ref"]["value"] == "FIXED-1"  # read-only fields are left alone
    assert after["city"]["value"] == "Kochi"
    assert [after[k]["value"] for k in radio] == [False, True]
    doc = pdf(path)
    parent = doc.xref_get_key(after[radio[1]]["xref"], "Parent")[1].split()[0]
    assert doc.xref_get_key(int(parent), "V")[1] == "/Medium"
    assert res["extra"]["filled"] == 7  # the read-only field and the 'off' radio do not count


def test_fill_form_flatten(run_tool, samples):
    fields = _fields(samples["form"])
    res, out = run_tool("fill_form", [samples["form"]], {"values": {str(fields["full_name"]["xref"]): "Grace Hopper"}, "flatten": True})
    doc = pdf(out / ok(res)["file"])
    assert not list(doc[0].widgets()) and "Grace Hopper" in doc[0].get_text()


def test_fill_form_without_fields(run_tool, samples):
    res, _ = run_tool("fill_form", [samples["pdf"]], {"values": {}})
    assert res["status"] == "error" and "no fillable" in res["message"].lower()


def test_detect_headings(samples):
    from pdfplume.structure import detect_headings
    heads = detect_headings(pdf(samples["pdf"]))
    assert [h["title"] for h in heads] == [f"Page {i} heading" for i in range(1, 6)]
    assert all(h["level"] == 1 for h in heads) and [h["page"] for h in heads] == list(range(5))


def test_normalize_toc():
    from pdfplume.structure import normalize_toc
    toc = normalize_toc([{"level": 2, "title": "A", "page": 1}, {"level": 4, "title": "B", "page": 99},
                         {"level": 1, "title": "  ", "page": 1}, {"level": 1, "title": "C", "page": 0}], 5)
    assert toc == [[1, "A", 1], [2, "B", 5], [1, "C", 1]]


def test_edit_metadata(run_tool, samples):
    import pikepdf
    toc = [{"level": 1, "title": "Intro", "page": 1}, {"level": 2, "title": "Details", "page": 2},
           {"level": 1, "title": "End", "page": 5}]
    res, out = run_tool("edit_metadata", [samples["pdf"]], {
        "metadata": {"title": "Annual report", "author": "Visu", "keywords": "pdf, report"}, "toc": toc})
    path = out / ok(res)["file"]
    doc = pdf(path)
    assert doc.metadata["title"] == "Annual report" and doc.metadata["author"] == "Visu"
    assert doc.metadata["subject"] == ""
    assert doc.get_toc() == [[1, "Intro", 1], [2, "Details", 2], [1, "End", 5]]
    assert path.name == "sample.pdf" and res["extra"]["bookmarks"] == 3
    with pikepdf.open(path) as p:
        meta = p.open_metadata()
        assert meta.get("dc:title") == "Annual report"


def test_edit_metadata_keeps_bookmarks_when_not_given(run_tool, samples, tmp_path):
    src = pdf(samples["pdf"])
    src.set_toc([[1, "Keep me", 3]])
    src.save(tmp_path / "with_toc.pdf")
    res, out = run_tool("edit_metadata", [tmp_path / "with_toc.pdf"], {"metadata": {"title": "X"}})
    assert pdf(out / ok(res)["file"]).get_toc() == [[1, "Keep me", 3]]


# ------------------------------------------------------------------- chat
def _index(path):
    from pdfplume import chat
    return chat.build_index(pdf(path))


def test_chat_index(samples):
    index = _index(samples["pdf"])
    assert index["pages"] == 5 and index["chunks"] and index["words"] > 200
    assert {c["page"] for c in index["chunks"]} == set(range(5))
    for c in index["chunks"]:
        x0, y0, x1, y1 = c["rect"]
        assert 0 <= x0 < x1 <= 1 and 0 <= y0 < y1 <= 1


def test_chat_offline_answer(samples):
    from pdfplume import chat
    index = _index(samples["pdf"])
    res = chat.answer(index, "What does compression do?")
    assert res["engine"] == "offline" and "ompression" in res["answer"] and "[p. " in res["answer"]
    assert res["sources"] and all(0 <= s["page"] < 5 for s in res["sources"])
    miss = chat.answer(index, "zebra giraffe safari")
    assert "couldn’t find" in miss["answer"] and miss["sources"] == []
    summ = chat.answer(index, "Can you summarize this document?")
    assert summ["answer"].startswith("•") and "[p. " in summ["answer"]


def test_chat_with_ollama(samples, monkeypatch):
    from pdfplume import chat, llm
    calls = []
    monkeypatch.setattr(llm, "status", lambda max_age=30: {"available": True, "model": "test", "url": "x", "models": ["test"]})
    monkeypatch.setattr(llm, "chat", lambda messages, timeout=600: calls.append(messages) or "Call +1 555 123 4567 [p. 1].")
    index = _index(samples["pdf"])
    history = [{"role": "user", "content": "Who do I contact?"}, {"role": "assistant", "content": "hello@example.com [p. 1]"}]
    res = chat.answer(index, "And the phone?", history)
    assert res["engine"] == "ollama" and res["model"] == "test" and res["answer"].endswith("[p. 1].")
    msgs = calls[0]
    assert msgs[0]["role"] == "system" and msgs[1:3] == history
    assert "[p. 1]" in msgs[-1]["content"] and "And the phone?" in msgs[-1]["content"]


def test_chat_ollama_failure_falls_back(samples, monkeypatch):
    from pdfplume import chat, llm
    monkeypatch.setattr(llm, "status", lambda max_age=30: {"available": True, "model": "m", "url": "x", "models": []})

    def boom(*a, **k):
        raise OSError("connection refused")
    monkeypatch.setattr(llm, "chat", boom)
    res = chat.answer(_index(samples["pdf"]), "What does merging do?")
    assert res["engine"] == "offline" and "warning" in res and "erging" in res["answer"]


# ------------------------------------------------------------- scan crop
def test_detect_document(samples):
    import numpy as np
    from PIL import Image
    from conftest import DESK_PAGE
    from pdfplume.scan import detect_document
    res = detect_document(Image.open(samples["desk"]))
    assert res["found"]
    got = np.array(res["corners"]) * [1599, 1199]
    assert np.abs(got - np.array(DESK_PAGE)).max() < 15
    blank = detect_document(Image.new("RGB", (800, 600), (90, 60, 40)))
    assert blank == {"corners": [[0.0, 0.0], [1.0, 0.0], [1.0, 1.0], [0.0, 1.0]], "found": False}


def test_warp_straightens_page(samples):
    from PIL import Image
    from pdfplume.scan import detect_document, warp
    im = Image.open(samples["desk"])
    out = warp(im, detect_document(im)["corners"])
    assert 0.7 < out.width / out.height < 0.9  # the sheet is portrait, about 850 x 1000 in the photo
    px = out.convert("L").resize((40, 50)).load()
    assert min(px[0, 0], px[39, 0], px[0, 49], px[39, 49]) > 120  # corners are paper, not desk


def _scan_page_ratio(res, out):
    page = pdf(out / ok(res)["file"])[0]
    xref = page.get_images()[0][0]
    info = page.parent.extract_image(xref)
    return info["width"] / info["height"]


def test_scan_auto_crop(run_tool, samples):
    res, out = run_tool("scan_to_pdf", [samples["desk"]], {"filter": "color"})
    assert 0.7 < _scan_page_ratio(res, out) < 0.9
    res, out = run_tool("scan_to_pdf", [samples["desk"]], {"filter": "color", "auto_crop": False})
    assert abs(_scan_page_ratio(res, out) - 4 / 3) < 0.01


def test_scan_manual_corners(run_tool, samples):
    corners = [[0.1, 0.1], [0.6, 0.1], [0.6, 0.9], [0.1, 0.9]]  # 800 x 960 px of the photo
    res, out = run_tool("scan_to_pdf", [samples["desk"]], {"filter": "color"}, file_options=[{"corners": corners}])
    assert abs(_scan_page_ratio(res, out) - 800 / 960) < 0.02
    res, out = run_tool("scan_to_pdf", [samples["desk"]], {"filter": "color"}, file_options=[{"crop": False}])
    assert abs(_scan_page_ratio(res, out) - 4 / 3) < 0.01


# --------------------------------------------------------- edit existing text
def test_text_lines(samples):
    from pdfplume.tools.edit import text_lines
    lines = text_lines(pdf(samples["pdf"])[0])
    assert lines[0]["text"] == "Page 1 heading" and lines[0]["size"] == 22 and lines[0]["font"] == "helvetica"
    assert lines[1]["text"].startswith("PDF tools make documents easier")
    assert all(0 <= l["rect"][0] < l["rect"][2] <= 1 for l in lines)
    assert text_lines(pdf(samples["rotated"])[0]) == []  # text runs vertically on the turned page


def test_edit_existing_text(run_tool, samples):
    from pdfplume.tools.edit import text_lines
    heading, first = text_lines(pdf(samples["pdf"])[0])[:2]
    elements = [
        {"type": "text", "page": 0, "x": heading["rect"][0], "y": heading["rect"][1], "w": 0.4, "h": 0.03,
         "text": "Quarterly results", "size": heading["size"], "color": heading["color"], "font": "helvetica", "replace": heading["rect"]},
        {"type": "text", "page": 0, "x": first["rect"][0], "y": first["rect"][1], "w": 0.1, "h": 0.02,
         "text": "", "size": 11, "replace": first["rect"]},
    ]
    res, out = run_tool("edit_pdf", [samples["pdf"]], {"elements": elements})
    page = pdf(out / ok(res)["file"])[0]
    text = page.get_text()
    assert "Quarterly results" in text and "Page 1 heading" not in text
    assert "PDF tools make documents easier to manage" not in text  # the deleted first line
    assert "separates pages into new documents" in text  # the next line is untouched
    assert page.get_images()  # the picture on the page survives
    hit = page.search_for("Quarterly results")[0]
    assert abs(hit.y0 - heading["rect"][1] * page.rect.height) < 6


# ------------------------------------------------------- digital signatures
def _png_data_url():
    import base64
    import io as _io
    from PIL import Image as _Image
    buf = _io.BytesIO()
    _Image.new("RGBA", (200, 80), (20, 20, 80, 255)).save(buf, "PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


def test_certificate_store(tmp_path):
    from pdfplume.certs import CertError, CertStore
    store = CertStore(tmp_path / "certs")
    assert store.list() == []
    info = store.create("Ada Lovelace", "pw1234", email="ada@example.com", organization="Analytical Engines")
    assert info["name"] == "Ada Lovelace" and info["self_signed"] and info["source"] == "created"
    assert len(info["fingerprint"]) == 64 and store.list()[0]["id"] == info["id"]
    assert b"BEGIN CERTIFICATE" in store.public_pem(info["id"]) and b"PRIVATE" not in store.public_pem(info["id"])
    with pytest.raises(CertError):
        store.create("Ada", "pw")  # password too short
    data = store.path(info["id"]).read_bytes()
    with pytest.raises(CertError):
        store.import_p12(data, "wrong")
    imported = store.import_p12(data, "pw1234")
    assert imported["source"] == "imported" and imported["fingerprint"] == info["fingerprint"]
    assert len(store.trusted()) == 2
    store.delete(info["id"])
    assert [c["id"] for c in store.list()] == [imported["id"]]
    with pytest.raises(CertError):
        store.path("../../etc/passwd")


def test_sign_digitally(run_tool, samples, tmp_path):
    from pdfplume.certs import CertStore, verify_pdf
    store = CertStore(tmp_path / "certs")
    cert = store.create("Grace Hopper", "secret99")
    el = {"type": "image", "page": 1, "x": 0.1, "y": 0.7, "w": 0.3, "h": 0.06, "src": _png_data_url()}
    digital = {"cert_path": str(store.path(cert["id"])), "password": "secret99", "reason": "Approved", "location": "Kochi"}
    res, out = run_tool("sign_pdf", [samples["pdf"]], {"elements": [el], "digital": digital})
    path = out / ok(res)["file"]
    sigs = res["extra"]["digital"]["signatures"]
    assert len(sigs) == 1 and sigs[0]["intact"] and sigs[0]["valid"] and sigs[0]["trusted"] and sigs[0]["covers_whole_file"]
    assert "Grace Hopper" in res["extra"]["digital"]["signer"]
    assert verify_pdf(path)[0]["trusted"] is False  # nobody else trusts a self-made certificate
    doc = pdf(path)
    widget = next(doc[1].widgets())
    assert widget.field_type == fitz.PDF_WIDGET_TYPE_SIGNATURE
    r, W, H = widget.rect, doc[1].rect.width, doc[1].rect.height
    assert abs(r.x0 - 0.1 * W) < 2 and abs(r.y0 - 0.7 * H) < 2  # the field sits over the signature image
    # Changing a byte after signing is detected.
    data = bytearray(path.read_bytes())
    i = data.find(b"Page 1 heading")
    if i < 0:
        i = len(data) // 3
    data[i] ^= 1
    tampered = tmp_path / "tampered.pdf"
    tampered.write_bytes(bytes(data))
    assert not all(s["intact"] for s in verify_pdf(tampered)) or not verify_pdf(tampered)


def test_sign_digitally_wrong_password(run_tool, samples, tmp_path):
    from pdfplume.certs import CertStore
    store = CertStore(tmp_path / "certs")
    cert = store.create("Grace Hopper", "secret99")
    res, _ = run_tool("sign_pdf", [samples["pdf"]], {"elements": [], "digital": {"cert_path": str(store.path(cert["id"])), "password": "nope"}})
    assert res["status"] == "error" and "password" in res["message"].lower()


def test_sign_digitally_twice(run_tool, samples, tmp_path):
    from pdfplume.certs import CertStore, verify_pdf
    store = CertStore(tmp_path / "certs")
    cert = store.create("Signer One", "secret99")
    digital = {"cert_path": str(store.path(cert["id"])), "password": "secret99"}
    res, out = run_tool("sign_pdf", [samples["pdf"]], {"digital": digital})
    first = out / ok(res)["file"]
    res2, out2 = run_tool("sign_pdf", [first], {"digital": digital})
    sigs = verify_pdf(out2 / ok(res2)["file"], trusted=store.trusted())
    assert [s["field"] for s in sigs] == ["Signature1", "Signature2"] and all(s["intact"] for s in sigs)


def test_visual_sign_keeps_digital_signature(run_tool, samples, tmp_path):
    from pdfplume.certs import CertStore, verify_pdf
    store = CertStore(tmp_path / "certs")
    cert = store.create("Signer One", "secret99")
    res, out = run_tool("sign_pdf", [samples["pdf"]], {"digital": {"cert_path": str(store.path(cert["id"])), "password": "secret99"}})
    signed = out / ok(res)["file"]
    el = {"type": "image", "page": 0, "x": 0.5, "y": 0.8, "w": 0.3, "h": 0.06, "src": _png_data_url()}
    res2, out2 = run_tool("sign_pdf", [signed], {"elements": [el]})
    sigs = verify_pdf(out2 / ok(res2)["file"], trusted=store.trusted())
    assert len(sigs) == 1 and sigs[0]["intact"] and not sigs[0]["covers_whole_file"]
    assert "stay valid" in res2["extra"]["note"]
